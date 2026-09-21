"""Сборка vacancies.xlsx: лист "Вакансии" + лист "Сводка"."""
from __future__ import annotations

import logging
from collections import Counter
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import EXCEL_FILE
from .geo import GEO_RU
from .models import Vacancy, sort_key

log = logging.getLogger(__name__)

SOURCE_NAMES = {
    "staff": "staff.am",
    "worknet": "worknet.am",
    "list": "list.am",
    "hire": "hire.am",
}

HEADERS = ["№", "Источник", "Вакансия (ориг.)", "Вакансия (рус.)", "Зарплата",
           "Локация", "Гео", "Компания", "Категория", "Удалённо", "Опубликовано", "Ссылка"]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
ALT_FILL = PatternFill("solid", fgColor="EDF2F8")
THIN = Side(style="thin", color="C9D4E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK_FONT = Font(color="0563C1", underline="single")


def _fmt_posted(v: Vacancy) -> str:
    if not v.posted_at:
        return ""
    try:
        dt = datetime.fromisoformat(v.posted_at.replace("Z", "+00:00"))
        return dt.strftime("%d.%m.%Y %H:%M")
    except ValueError:
        return v.posted_at


def build_excel(vacancies: list[Vacancy]) -> None:
    vacancies = sorted(vacancies, key=sort_key)
    wb = Workbook()
    ws = wb.active
    ws.title = "Вакансии"

    ws.append(HEADERS)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    for idx, v in enumerate(vacancies, 1):
        row = [
            idx,
            SOURCE_NAMES.get(v.source, v.source),
            v.title_orig,
            v.title_ru or ("" if v.title_orig else "—"),
            v.salary,
            v.city,
            GEO_RU.get(v.geo, v.geo),
            v.company,
            v.category,
            "да" if v.is_remote else "",
            _fmt_posted(v),
            v.url,
        ]
        ws.append(row)
        r = ws.max_row
        for c in ws[r]:
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=False)
        if idx % 2 == 0:
            for c in ws[r]:
                c.fill = ALT_FILL
        # гиперссылка
        url_cell = ws.cell(row=r, column=12)
        if v.url:
            url_cell.hyperlink = v.url
            url_cell.font = LINK_FONT
        if v.title_ru:
            ws.cell(row=r, column=4).font = Font(bold=True)

    # ширины колонок
    widths = [5, 11, 42, 42, 18, 18, 12, 26, 26, 9, 17, 46]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{ws.max_row}"

    # ---- Сводка ----
    st = wb.create_sheet("Сводка")
    by_source = Counter(SOURCE_NAMES.get(v.source, v.source) for v in vacancies)
    by_city = Counter(v.city or "не указана" for v in vacancies)
    remote_cnt = sum(1 for v in vacancies if v.is_remote)

    st.append(["Отчёт по вакансиям Армении", ""])
    st["A1"].font = Font(bold=True, size=14)
    st.append(["Сформировано", datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M UTC")])
    st.append(["Всего вакансий", len(vacancies)])
    st.append(["Из них зарубежных (скрыты фильтром сайта)",
               sum(1 for v in vacancies if v.geo == "foreign")])
    st.append(["Удалённых", remote_cnt])
    st.append([])
    st.append(["По источникам", ""])
    st.cell(row=st.max_row, column=1).font = Font(bold=True)
    for name, cnt in by_source.most_common():
        st.append([name, cnt])
    st.append([])
    st.append(["Топ локаций", ""])
    st.cell(row=st.max_row, column=1).font = Font(bold=True)
    for name, cnt in by_city.most_common(20):
        st.append([name, cnt])
    st.column_dimensions["A"].width = 32
    st.column_dimensions["B"].width = 12

    EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_FILE)
    log.info("Excel сохранён: %s (%d вакансий)", EXCEL_FILE, len(vacancies))
